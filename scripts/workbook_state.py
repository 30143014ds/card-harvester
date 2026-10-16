"""Restore working state from the delivered XLSX, or start an isolated new table."""
from __future__ import annotations

import argparse
import copy
import io
import json
import re
import shutil
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from business_cards.contacts import (VISIBLE_FIELDS,
                                     atomic_json, gallery_cell, load_catalog, save_catalog)
from business_cards.schema import SUPPORTED_HEADERS, FIELD_ALIASES, get_sheet


def new_table(directory: Path):
    directory = directory.resolve()
    # A fresh target is essential: "new table" must never reset the cumulative one.
    directory.mkdir(parents=True, exist_ok=False)
    library = directory / 'card_library'
    library.mkdir()
    atomic_json(library / 'catalog.json', {'version': 2, 'sources': {}, 'cards': []})
    return {'library': str(library), 'xlsx': str(directory / 'business_cards.xlsx')}


def restore(workbook_path: Path, library: Path):
    workbook_path, library = workbook_path.resolve(), library.resolve()
    book = load_workbook(workbook_path)
    try:
        sheet = get_sheet(book, 'contacts')
        headers = [cell.value for cell in sheet[1]]
        if headers not in SUPPORTED_HEADERS:
            raise ValueError('Expected a business-card workbook with hidden _id')
        id_col = headers.index('_id') + 1
        rows, evidence = {}, {}
        for row in sheet.iter_rows(min_row=2):
            stable_id = row[id_col - 1].value
            if not stable_id and not any(cell.value for cell in row):
                continue
            if not isinstance(stable_id, str) or not re.fullmatch(r'c_[0-9a-f]{24}', stable_id) or stable_id in rows:
                raise ValueError('Missing/duplicate/invalid hidden contact ID')
            values = {}
            for name, cell in zip(VISIBLE_FIELDS, row[:5]):
                text = str(cell.value).strip() if cell.value is not None else ''
                values[name] = [v.strip() for v in re.split(r'[\n;]+', text) if v.strip()] if name in ('emails', 'phones') else text or None
            rows[stable_id] = values
            if row[0].comment:
                match = re.search(r'(?:Company from domain|Компания по домену) (\S+): (https?://\S+)', row[0].comment.text)
                if match:
                    evidence[stable_id] = {'domain': match[1], 'url': match[2], 'method': 'restored_workbook_comment'}
        if not rows:
            raise ValueError('Workbook contains no contacts')

        origins = {stable_id: [] for stable_id in rows}
        source_order, sources = [], {}
        source_sheet = get_sheet(book, 'sources', required=False)
        if source_sheet is not None:
            for stable_id, video, card_id, timestamp, digest, *_ in source_sheet.iter_rows(min_row=2, values_only=True):
                if stable_id not in rows:
                    continue
                if not isinstance(digest, str) or not re.fullmatch(r'[0-9a-f]{64}', digest):
                    raise ValueError('Invalid source SHA256 in workbook')
                origins[stable_id].append({'video_sha256': digest, 'video': video, 'card_id': card_id, 'time_sec': timestamp})
                if stable_id not in source_order:
                    source_order.append(stable_id)
                source = sources.setdefault(digest, {'video': video, 'source_file': None, 'restored_from_xlsx': True, 'source_materials_retained': False, 'cards': 0})
                source['cards'] += 1

        issues = {stable_id: [] for stable_id in rows}
        review_sheet = get_sheet(book, 'review', required=False)
        if review_sheet is not None:
            for stable_id, _, field, fragment, reason, *_ in review_sheet.iter_rows(min_row=2, values_only=True):
                if stable_id in issues:
                    issues[stable_id].append({'field': FIELD_ALIASES.get(field, field), 'visible_fragment': fragment or '', 'reason': reason or ''})

        existing = (library / 'catalog.json').is_file()
        catalog = copy.deepcopy(load_catalog(library))
        if existing and catalog['cards']:
            if {c['id'] for c in catalog['cards']} != set(rows):
                raise ValueError('Workbook and working catalog contain different contacts; use a fresh library to restore this workbook')
            ordered_ids = [c['id'] for c in catalog['cards']]
        else:
            ordered_ids = source_order + [stable_id for stable_id in rows if stable_id not in source_order]

        photo_data, gallery_order = {}, []
        gallery = get_sheet(book, 'originals', required=False)
        if gallery is not None:
            expected = {gallery_cell(n): stable_id for n, stable_id in enumerate(ordered_ids)}
            for picture in sorted(gallery._images, key=lambda p: (p.anchor._from.row, p.anchor._from.col)):
                anchor = picture.anchor._from
                cell = gallery.cell(anchor.row + 1, anchor.col + 1)
                stable_id = cell.comment.text.strip() if cell.comment else expected.get((anchor.row, anchor.col))
                if stable_id not in rows or stable_id in photo_data:
                    raise ValueError('Gallery image cannot be matched to a unique contact ID')
                photo_data[stable_id] = picture._data()
                gallery_order.append(stable_id)
        else:
            # Previous workbook format: table photos anchored to the contact row.
            for picture in sheet._images:
                stable_id = sheet.cell(picture.anchor._from.row + 1, id_col).value
                if stable_id not in rows or stable_id in photo_data:
                    raise ValueError('Table image cannot be matched to a unique contact ID')
                photo_data[stable_id] = picture._data()
        if set(photo_data) != set(rows):
            raise ValueError('Workbook is missing contact photos; no catalog was replaced')
        if not existing or not catalog['cards']:
            ordered_ids = gallery_order or ordered_ids
            cards = []
            for stable_id in ordered_ids:
                values = rows[stable_id]
                cards.append({'id': stable_id, **values, 'company_evidence': evidence.get(stable_id),
                              'photo': '', 'raw': {'restored_from_xlsx': True},
                              'issues': issues[stable_id], 'origins': origins[stable_id],
                              'manual': copy.deepcopy(values)})
            catalog = {'version': 2, 'sources': sources, 'cards': cards}

        library.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix='restore-', dir=library) as temporary:
            staging = Path(temporary)
            for card in catalog['cards']:
                data = photo_data[card['id']]
                with Image.open(io.BytesIO(data)) as photo:
                    photo.verify()
                    extension = 'jpg' if photo.format == 'JPEG' else 'png'
                relative = f'photos/{card["id"]}.{extension}'
                (staging / 'photos').mkdir(exist_ok=True)
                (staging / relative).write_bytes(data)
                card['photo'] = relative
            (library / 'photos').mkdir(exist_ok=True)
            for image_path in (staging / 'photos').iterdir():
                shutil.copy2(image_path, library / 'photos' / image_path.name)
            save_catalog(library, catalog)
        return {'cards': len(rows), 'photos': len(photo_data), 'library': str(library), 'workbook': str(workbook_path)}
    finally:
        book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    fresh = commands.add_parser('new', help='Create an isolated empty library; never overwrite a folder')
    fresh.add_argument('directory', type=Path)
    recover = commands.add_parser('restore', help='Recover contact IDs, data, sources, review flags and photos from XLSX')
    recover.add_argument('xlsx', type=Path)
    recover.add_argument('--library', type=Path, required=True)
    args = parser.parse_args()
    result = new_table(args.directory) if args.command == 'new' else restore(args.xlsx, args.library)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
