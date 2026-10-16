import hashlib,json,tempfile,unittest
from pathlib import Path
from zipfile import ZipFile
from PIL import Image
from openpyxl import load_workbook
from business_cards.contacts import atomic_json,export_library,import_batch,import_edits,load_catalog
from business_cards.schema import RUSSIAN_HEADERS
from workbook_state import new_table,restore
from extract_cards import apply_review_choices

class WorkbookStateTests(unittest.TestCase):
 def test_legacy_labels_restore_photos_evidence_flags_and_manual_edits(self):
  with tempfile.TemporaryDirectory() as d:
   base=Path(d);library=base/'library';batch=self.batch(base,'a','Alice')
   recognition=batch/'recognition/card_001.json'
   record=json.loads(recognition.read_text())
   record.update(company_evidence={'domain':'acme.test','url':'https://acme.test/'},
                 uncertain_fields=[{'field':'position','visible_fragment':'Dir','reason':'Glare'}])
   atomic_json(recognition,record);import_batch(library,batch,{})
   xlsx=base/'legacy.xlsx';export_library(library,xlsx)
   book=load_workbook(xlsx)
   for english,russian in [('Contacts','Визитки'),('Originals','Оригиналы'),('Review','Проверить'),('Sources','Источники')]:
    book[english].title=russian
   sheet=book['Визитки']
   for cell,label in zip(sheet[1],RUSSIAN_HEADERS):cell.value=label
   sheet['B2']='Alice edited'
   sheet['A2'].comment.text='Компания по домену acme.test: https://acme.test/'
   book['Проверить']['C2']='Должность'
   book.save(xlsx);book.close()
   self.assertEqual(import_edits(library,xlsx)['manual_fields_saved'],1)
   recovered=base/'restored';restore(xlsx,recovered)
   card=load_catalog(recovered)['cards'][0]
   self.assertEqual(card['full_name'],'Alice edited')
   self.assertEqual(card['company_evidence']['domain'],'acme.test')
   self.assertEqual(card['issues'][0]['field'],'position')
   self.assertTrue((recovered/card['photo']).is_file())
   output=base/'english.xlsx';export_library(recovered,output)
   book=load_workbook(output)
   self.assertEqual(book.sheetnames,['Contacts','Originals','Review','Sources'])
   self.assertEqual(book['Contacts']['G2'].hyperlink.location,"'Originals'!A4")
   book.close()
 def test_review_choices_ignore_unselected_scenes(self):
  with tempfile.TemporaryDirectory() as d:
   root=Path(d)
   atomic_json(root/'review_choices.json',{'card_001':{'selected_frame':12},'card_002':{'selected_frame':36}})
   ranked=apply_review_choices(root,{'card_001':[{'frame':12}],'card_002':[]})
   self.assertEqual(ranked['card_001'][0]['frame'],12)
 def batch(self,base,digest,name):
  p=base/digest;(p/'cards').mkdir(parents=True);(p/'recognition').mkdir()
  Image.new('RGB',(200,120),'red' if name=='Alice' else 'blue').save(p/'cards/card_001.png')
  atomic_json(p/'manifest.json',{'source_sha256':digest*64,'source_video':digest+'.mov','count':1,'cards':[{'id':'card_001','selected':{'file':'cards/card_001.png','time_sec':1.0}}]})
  atomic_json(p/'recognition/card_001.json',{'id':'card_001','company':'Acme','full_name':name,'position':'Director','emails':[name.lower()+'@acme.test'],'phones':['+7 495 123-45-67, доб. 1480'],'uncertain_fields':[],'raw_text':''})
  return p
 def test_restore_sorted_table_keeps_gallery_ids_photos_edits_then_appends(self):
  with tempfile.TemporaryDirectory() as d:
   base=Path(d);original=base/'original';a=self.batch(base,'a','Alice');b=self.batch(base,'b','Bob')
   import_batch(original,a,{});import_batch(original,b,{})
   ids=[c['id'] for c in load_catalog(original)['cards']]
   xlsx=base/'out.xlsx';export_library(original,xlsx)
   w=load_workbook(xlsx);s=w['Contacts'];data=[list(row) for row in s.iter_rows(min_row=2,values_only=True)]
   for row,values in zip([2,3],reversed(data)):
    for col in [1,2,3,4,5,8]:s.cell(row,col,values[col-1])
   s['B2']='Bob edited';w.save(xlsx);w.close()
   with ZipFile(xlsx) as z:before={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('xl/media/')}
   target=base/'recovered';restore(xlsx,target);c=load_catalog(target)
   self.assertEqual([c['id'] for c in c['cards']],ids)
   self.assertEqual([c['full_name'] for c in c['cards']],['Alice','Bob edited'])
   self.assertIn('ext. 1480',c['cards'][0]['phones'][0])
   recovered=base/'recovered.xlsx';export_library(target,recovered)
   with ZipFile(recovered) as z:after={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('xl/media/')}
   self.assertEqual(before,after)
   result=import_batch(target,self.batch(base,'c','Charlie'),{})
   self.assertEqual(result['total'],3)
   self.assertEqual(load_catalog(target)['cards'][1]['full_name'],'Bob edited')
   export_library(target,recovered)
   w=load_workbook(recovered);self.assertEqual(len(w['Originals']._images),3);w.close()
 def test_new_table_is_isolated_and_refuses_overwrite(self):
  with tempfile.TemporaryDirectory() as d:
   base=Path(d);current=base/'current.xlsx';current.write_bytes(b'keep current')
   result=new_table(base/'new');self.assertEqual(load_catalog(Path(result['library']))['cards'],[])
   with self.assertRaises(FileExistsError):new_table(base/'new')
   self.assertEqual(current.read_bytes(),b'keep current')

if __name__=='__main__':unittest.main()
