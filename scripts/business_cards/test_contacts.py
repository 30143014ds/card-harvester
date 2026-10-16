import json
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from openpyxl import Workbook, load_workbook

from business_cards.contacts import (HEADERS, atomic_json, collect_phones, export_library,
                                     import_batch, import_edits, load_catalog,
                                     normalize_phone, resolve_company, gallery_cell)


class ContactLibraryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.base = Path(self.temp.name)
        self.root = self.base / 'library'

    def tearDown(self):
        self.temp.cleanup()

    def batch(self, name, digest, position='Director'):
        path = self.base / name
        (path / 'cards').mkdir(parents=True)
        (path / 'recognition').mkdir()
        Image.new('RGB', (160, 90), 'white').save(path / 'cards/card_001.png')
        atomic_json(path / 'manifest.json', {
            'source_sha256': digest * 64, 'source_video': name + '.mov', 'count': 1,
            'cards': [{'id': 'card_001', 'selected': {'file': 'cards/card_001.png', 'time_sec': 1.0}}]})
        atomic_json(path / 'recognition/card_001.json', {
            'id': 'card_001', 'company': 'Acme', 'full_name': 'Alex', 'position': position,
            'emails': ['alex@acme.test'], 'phones': ['+7 (916) 123 45 67'],
            'uncertain_fields': [], 'raw_text': ''})
        return path

    def test_phone_merge_format_and_extension(self):
        record = {'work_phones': ['+7 (495) 336-16-19', '113'],
                  'mobile_phones': ['8 916 633 5891'],
                  'unclassified_phones': ['+7 (916) 633-58-91'],
                  'fax_phones': ['+7 495 999 99 99'], 'raw_text': 'внутр. тел.: 113'}
        self.assertEqual(collect_phones(record), ['+7 495 336-16-19, ext. 113', '8 916 633 5891', '+7 916 633-58-91'])
        self.assertEqual(normalize_phone('+49 (0)89 2180-6979'), '+49 89 2180-69-79')
        self.assertEqual(normalize_phone('0212019100'), '021 201 91 00')
        self.assertEqual(normalize_phone('237-327'), '237-327')
        self.assertEqual(normalize_phone('+7 495 212 21 56 ext. 435'), '+7 495 212-21-56, ext. 435')
        self.assertEqual(normalize_phone('+7 495 419-21-46, доб. 79 (59-35)'),
                         '+7 495 419-21-46, ext. 79 (59-35)')

    def test_phone_country_is_not_inferred_from_language_or_other_numbers(self):
        record = {'phones': ['+7 916 000 00 01', '89160000001', '9160000001'],
                  'raw_text': 'Компания'}
        values = collect_phones(record)
        self.assertEqual(len(values), 3)
        self.assertTrue(values[0].startswith('+7'))
        self.assertFalse(values[1].startswith('+'))
        self.assertFalse(values[2].startswith('+'))
        self.assertEqual(collect_phones({'phones': ['1234 ext. 56', '123456']}),
                         ['1234, ext. 56', '123456'])

    def test_reviewed_back_is_audited_without_a_duplicate_contact(self):
        batch = self.batch('back', 'a')
        path = batch / 'recognition/card_001.json'
        record = json.loads(path.read_text())
        record.update(skip_as_contact=True, skip_reason='QR repeats the reviewed front')
        atomic_json(path, record)
        self.assertEqual(import_batch(self.root, batch, {})['non_contact_sides'], 1)
        catalog = load_catalog(self.root)
        self.assertEqual(catalog['cards'], [])
        self.assertEqual(catalog['sources']['a' * 64]['skipped_sides'][0]['card_id'], 'card_001')

    def test_domain_lookup_does_not_invent_employer_or_mailbox(self):
        known = {'acme.test': {'company': 'Acme', 'url': 'https://acme.test/'}}
        self.assertEqual(resolve_company({'emails': ['alex@gmail.com']}, known), (None, None))
        company, proof = resolve_company({'emails': [], 'uncertain_fields': [
            {'field': 'emails', 'visible_fragment': '…@acme.test'}]}, known)
        self.assertEqual(company, 'Acme')
        self.assertEqual(proof['domain'], 'acme.test')

    def test_repeat_import_and_exact_duplicate_and_new_role(self):
        first = self.batch('first', 'a')
        self.assertEqual(import_batch(self.root, first, {})['total'], 1)
        stable = load_catalog(self.root)['cards'][0]['id']
        self.assertEqual(import_batch(self.root, first, {})['added'], 0)
        self.assertEqual(load_catalog(self.root)['cards'][0]['id'], stable)
        identical = self.batch('second', 'b')
        self.assertEqual(import_batch(self.root, identical, {})['duplicates_merged'], 1)
        other_role = self.batch('third', 'c', position='Manager')
        self.assertEqual(import_batch(self.root, other_role, {})['total'], 2)
        self.assertEqual(len(load_catalog(self.root)['cards'][0]['origins']), 2)

    def test_visually_reviewed_duplicate_keeps_existing_row_and_manual_fields(self):
        first = self.batch('first', 'a')
        import_batch(self.root, first, {})
        stable = load_catalog(self.root)['cards'][0]['id']
        second = self.batch('second', 'b', position='Director; OCR variant')
        path = second / 'recognition/card_001.json'
        record = json.loads(path.read_text())
        record.update(duplicate_of=stable, duplicate_reason='Both photos reviewed; same printed card')
        atomic_json(path, record)
        for _ in range(2):
            self.assertEqual(import_batch(self.root, second, {})['total'], 1)
        catalog = load_catalog(self.root)
        self.assertEqual(catalog['cards'][0]['position'], 'Director')
        self.assertEqual(len(catalog['cards'][0]['origins']), 2)

    def test_manual_edits_follow_ids_and_survive_ocr_refresh(self):
        batch = self.batch('first', 'a')
        import_batch(self.root, batch, {})
        card = load_catalog(self.root)['cards'][0]
        file = self.base / 'edited.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.title = 'Contacts'
        sheet.append(HEADERS)
        sheet.append(['Acme', 'Alex corrected', 'Director', 'alex@acme.test',
                      '+7 916 123-45-67', '', '', card['id']])
        book.save(file)
        self.assertEqual(import_edits(self.root, file)['manual_fields_saved'], 1)
        import_batch(self.root, batch, {})
        self.assertEqual(load_catalog(self.root)['cards'][0]['full_name'], 'Alex corrected')

    def test_floating_images_link_to_fixed_gallery_and_share_compressed_media(self):
        import_batch(self.root, self.batch('first', 'a'), {})
        file = self.base / 'out.xlsx'
        result = export_library(self.root, file)
        self.assertEqual(result['cards'], 1)
        with ZipFile(file) as z:
            self.assertNotIn('xl/richData/rdrichvalue.xml', z.namelist())
            self.assertIn('xl/drawings/drawing1.xml', z.namelist())
            self.assertIn('xl/drawings/drawing2.xml', z.namelist())
            media = [name for name in z.namelist() if name.startswith('xl/media/')]
            self.assertEqual(len(media), 1)
            self.assertTrue(media[0].endswith('.jpeg'))
        workbook = load_workbook(file)
        self.assertEqual(workbook.sheetnames[:2], ['Contacts', 'Originals'])
        self.assertEqual(workbook['Contacts']['G2'].hyperlink.location, "'Originals'!A4")
        self.assertEqual(len(workbook['Contacts']._images), 1)
        self.assertEqual(len(workbook['Originals']._images), 1)
        self.assertEqual(workbook['Originals']['A3'].value, '#001 · Alex')
        self.assertIsNone(workbook['Originals'].auto_filter.ref)
        self.assertEqual([gallery_cell(i) for i in range(5)], [(3, 0), (3, 2), (3, 4), (3, 6), (6, 0)])

    def test_incomplete_batch_cannot_modify_existing_catalog(self):
        first = self.batch('first', 'a')
        import_batch(self.root, first, {})
        before = (self.root / 'catalog.json').read_bytes()
        broken = self.batch('broken', 'b')
        (broken / 'recognition/card_001.json').unlink()
        with self.assertRaises(FileNotFoundError):
            import_batch(self.root, broken, {})
        self.assertEqual((self.root / 'catalog.json').read_bytes(), before)


if __name__ == '__main__':
    unittest.main()
