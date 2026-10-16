"""Persistent contact library; exports are rebuilt from data, never from drawings."""
from __future__ import annotations

import hashlib
import io
import json
import re
import shutil
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

import phonenumbers
import xlsxwriter
from openpyxl import load_workbook
from PIL import Image
from xlsxwriter.utility import xl_rowcol_to_cell
from .schema import HEADERS, LEGACY_HEADERS, SUPPORTED_HEADERS, get_sheet

VISIBLE_FIELDS = ['company', 'full_name', 'position', 'emails', 'phones']
PERSONAL_DOMAINS = {'mail.ru', 'bk.ru', 'list.ru', 'inbox.ru', 'gmail.com', 'yandex.ru',
                    'ya.ru', 'outlook.com', 'hotmail.com', 'yahoo.com', 'yahoo.co.jp',
                    'icloud.com', 'aol.com', 'proton.me', 'protonmail.com'}
EXTENSION = re.compile(r'\s*\(?\b(?:доб(?:авочный)?\.?|ext\.?|ex\.?|add\.?|int\.?|вн(?:утр)?\.?)\s*[:.]?\s*([\d -]+(?:\([\d -]+\))?)\)?\s*$', re.I)
FIELD_NAMES = {'company': 'Company', 'full_name': 'Full name', 'position': 'Position',
               'emails': 'Email', 'phones': 'Phone', 'work_phones': 'Phone',
               'mobile_phones': 'Phone', 'unclassified_phones': 'Phone'}


def atomic_json(path: Path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + '.tmp')
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    temporary.replace(path)


def load_catalog(root: Path):
    path = root / 'catalog.json'
    if path.exists():
        data = json.loads(path.read_text(encoding='utf-8'))
        if data.get('version') != 2:
            raise ValueError('Unsupported catalog version')
        return data
    return {'version': 2, 'sources': {}, 'cards': []}


def save_catalog(root: Path, data):
    path = root / 'catalog.json'
    if path.exists():
        backups = root / 'backups'
        backups.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%f')
        shutil.copy2(path, backups / f'catalog_{stamp}.json')
    atomic_json(path, data)


def normalize_phone(value: str):
    text = unicodedata.normalize('NFKC', value).strip()
    match = EXTENSION.search(text)
    extension = match.group(1).strip() if match else None
    base = text[:match.start()].strip(' (,;/') if match else text
    # (0) is an optional national trunk prefix, not part of an international number.
    if base.startswith('+'):
        base = base.replace('(0)', '')
    digits = re.sub(r'\D', '', base)
    try:
        # A national number does not identify its country; preserve it without guessing.
        number = phonenumbers.parse(base, None)
        if not phonenumbers.is_possible_number(number):
            raise ValueError('Not enough digits for an international number')
        formatted = phonenumbers.format_number(number, phonenumbers.PhoneNumberFormat.INTERNATIONAL)
    except (phonenumbers.NumberParseException, ValueError):
        # Keep short/local numbers and unknown country codes; don't invent missing digits.
        formatted = re.sub(r'\s+', ' ', re.sub(r'[().]', ' ', base)).strip()
    if formatted.startswith('+'):
        # Some numbering plans leave a 7–9 digit subscriber number ungrouped.
        parts = formatted.split(' ')
        last = parts[-1]
        if last.isdigit() and len(last) >= 7:
            first = 3 if len(last) % 2 else 4
            parts[-1] = '-'.join([last[:first]] + [last[i:i + 2] for i in range(first, len(last), 2)])
            formatted = ' '.join(parts)
    elif formatted.isdigit() and len(formatted) in (7, 8, 9, 10):
        sizes = {7: [3, 2, 2], 8: [4, 4], 9: [3, 3, 3], 10: [3, 3, 2, 2]}[len(formatted)]
        offset, parts = 0, []
        for length in sizes:
            parts.append(formatted[offset:offset + length])
            offset += length
        formatted = ' '.join(parts)
    return formatted + (f', ext. {extension}' if extension else '')


def collect_phones(record):
    groups = record.get('phones')
    if groups is None:
        work = list(record.get('work_phones', []))
        short = [v for v in work if len(re.sub(r'\D', '', v)) < 6]
        full = [v for v in work if len(re.sub(r'\D', '', v)) >= 6]
        if len(full) == 1 and short and re.search(r'внутр|вн\.|ext|доб\.', record.get('raw_text', ''), re.I):
            work = [full[0] + ', ext. ' + ', '.join(short)]
        groups = work + record.get('mobile_phones', []) + record.get('unclassified_phones', [])
    if not isinstance(groups, list) or not all(isinstance(v, str) for v in groups):
        raise TypeError('phones must be a list of strings')
    result, seen = [], set()
    for value in groups:
        formatted = normalize_phone(value)
        match = EXTENSION.search(formatted)
        base = formatted[:match.start()] if match else formatted
        key = (re.sub(r'[^\d+]', '', base), re.sub(r'\D', '', match.group(1)) if match else '')
        if key[0] and key not in seen:
            result.append(formatted)
            seen.add(key)
    return result


def domain_candidates(record):
    domains = []
    for email in record.get('emails', []):
        if '@' in email:
            domains.append(email.rsplit('@', 1)[1].strip().lower())
    # A legible domain can identify a company even if the mailbox name is unreadable.
    for flag in record.get('uncertain_fields', []):
        if flag.get('field') == 'emails':
            domains.extend(re.findall(r'@([a-zA-Z0-9-]+(?:\.[a-zA-Z0-9-]+)+)', flag.get('visible_fragment', '')))
    return list(dict.fromkeys(d.lower() for d in domains if d.lower() not in PERSONAL_DOMAINS))


def resolve_company(record, domains):
    if record.get('company'):
        return record['company'], record.get('company_evidence')
    for domain in domain_candidates(record):
        info = domains.get(domain)
        if info:
            return info['company'], dict(info, domain=domain, method='email_domain')
    return None, None


def clean_key(value):
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFKC', value or '')).casefold().strip()


def fingerprint(card):
    # Same person in a different role is a separate card. Only exact structured duplicates merge.
    if not card.get('full_name') or not (card.get('emails') or card.get('phones')):
        return None
    values = [clean_key(card.get(key)) for key in ('company', 'full_name', 'position')]
    values += [sorted(clean_key(v) for v in card.get('emails', [])),
               sorted(re.sub(r'[^\d+]', '', v) for v in card.get('phones', []))]
    return hashlib.sha256(json.dumps(values, ensure_ascii=False).encode()).hexdigest()


def import_batch(root: Path, batch: Path, domains: dict):
    manifest = json.loads((batch / 'manifest.json').read_text())
    digest = manifest['source_sha256']
    if not re.fullmatch(r'[0-9a-f]{64}', digest):
        raise ValueError('Invalid video SHA256')
    if manifest['count'] != len(manifest['cards']):
        raise ValueError('Manifest count mismatch')
    catalog = load_catalog(root)
    known_ids = {card['id'] for card in catalog['cards']}
    prepared, local_ids, skipped = [], set(), []
    for source_card in manifest['cards']:
        local_id = source_card['id']
        if local_id in local_ids or not re.fullmatch(r'card_\d+', local_id):
            raise ValueError('Duplicate or invalid card ID')
        local_ids.add(local_id)
        record = json.loads((batch / 'recognition' / f'{local_id}.json').read_text())
        if record['id'] != local_id:
            raise ValueError('Recognition ID mismatch')
        duplicate_of = record.get('duplicate_of')
        if duplicate_of and (duplicate_of not in known_ids or not record.get('duplicate_reason')):
            raise ValueError(f'{local_id}: reviewed duplicate needs a known target ID and evidence')
        for key in ('company', 'full_name', 'position'):
            if record.get(key) is not None and not isinstance(record[key], str):
                raise TypeError(f'{local_id}.{key}: expected string or null')
        if not isinstance(record.get('emails'), list) or not all(isinstance(v, str) for v in record['emails']):
            raise TypeError(f'{local_id}.emails: expected list of strings')
        photo = (batch / source_card['selected']['file']).resolve()
        if not photo.is_relative_to(batch.resolve()) or not photo.is_file():
            raise ValueError(f'Missing/invalid photo path for {local_id}')
        if record.get('skip_as_contact'):
            reason = record.get('skip_reason')
            if not isinstance(reason, str) or not reason.strip():
                raise ValueError(f'{local_id}: skipped side needs a review reason')
            skipped.append({'card_id': local_id, 'reason': reason,
                            'related_card_id': record.get('related_card_id'),
                            'photo': str(photo), 'recognition': str((batch / 'recognition' / f'{local_id}.json').resolve())})
            continue
        stable_id = 'c_' + hashlib.sha256(f'{digest}:{local_id}'.encode()).hexdigest()[:24]
        company, evidence = resolve_company(record, domains)
        card = {'id': stable_id, 'company': company, 'full_name': record.get('full_name'),
                'position': record.get('position'), 'emails': record['emails'],
                'phones': collect_phones(record), 'company_evidence': evidence,
                'photo': f'photos/{stable_id}.png', 'raw': record,
                'issues': record.get('uncertain_fields', []),
                'origins': [{'video_sha256': digest, 'video': manifest['source_video'],
                             'card_id': local_id, 'time_sec': source_card['selected']['time_sec']}],
                'manual': {}}
        prepared.append((card, photo))
    by_origin = {(origin['video_sha256'], origin['card_id']): card for card in catalog['cards']
                 for origin in card['origins']}
    by_fingerprint = {fingerprint(card): card for card in catalog['cards'] if fingerprint(card)}
    by_id = {card['id']: card for card in catalog['cards']}
    root.mkdir(parents=True, exist_ok=True)
    (root / 'photos').mkdir(exist_ok=True)
    added = updated = merged = 0
    for card, photo in prepared:
        origin = card['origins'][0]
        existing = by_origin.get((digest, origin['card_id']))
        key = fingerprint(card)
        duplicate_of = card['raw'].get('duplicate_of')
        if duplicate_of:
            target = by_id[duplicate_of]
            if existing and existing is not target:
                raise ValueError('Reviewed duplicate conflicts with an already imported contact')
            if not existing:
                target['origins'].append(origin)
                by_origin[(digest, origin['card_id'])] = target
            merged += 1
            continue
        elif existing:
            # Refresh OCR while keeping IDs, other video occurrences and user edits.
            stable_id, origins, manual = existing['id'], existing['origins'], existing.get('manual', {})
            card.update(id=stable_id, origins=origins, manual=manual, photo=f'photos/{stable_id}.png')
            for field, value in manual.items():
                card[field] = value
            if 'company' in manual:
                card['company_evidence'] = None
            existing.clear()
            existing.update(card)
            updated += 1
        elif key and key in by_fingerprint:
            existing = by_fingerprint[key]
            existing['origins'].append(origin)
            by_origin[(digest, origin['card_id'])] = existing
            merged += 1
            continue
        else:
            catalog['cards'].append(card)
            existing = card
            by_origin[(digest, origin['card_id'])] = card
            if key:
                by_fingerprint[key] = card
            added += 1
        destination = root / existing['photo']
        if photo != destination.resolve():
            shutil.copy2(photo, destination)
    info_path = batch / 'video_info.json'
    info = json.loads(info_path.read_text()) if info_path.exists() else {}
    catalog['sources'][digest] = {'video': manifest['source_video'], 'cards': manifest['count'],
                                   'batch': str(batch.resolve()),
                                   'source_file': manifest.get('source_path', info.get('source')),
                                   'skipped_sides': skipped}
    save_catalog(root, catalog)
    return {'added': added, 'refreshed': updated, 'duplicates_merged': merged, 'non_contact_sides': len(skipped),
            'total': len(catalog['cards'])}


def import_edits(root: Path, workbook_path: Path):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = get_sheet(workbook, 'contacts')
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    if headers not in SUPPORTED_HEADERS:
        raise ValueError('Expected the contact workbook with hidden _id column')
    id_column = headers.index('_id')
    catalog = load_catalog(root)
    cards = {card['id']: card for card in catalog['cards']}
    seen, changes = set(), 0
    for values in sheet.iter_rows(min_row=2, max_col=len(headers), values_only=True):
        if not any(values):
            continue
        card_id = values[id_column]
        if card_id not in cards or card_id in seen:
            raise ValueError(f'Unknown, missing or duplicate _id: {card_id}')
        seen.add(card_id)
        card = cards[card_id]
        for field, value in zip(VISIBLE_FIELDS, values[:5]):
            value = str(value).strip() if value is not None else ''
            if field in ('emails', 'phones'):
                parsed = [v.strip() for v in re.split(r'[\n;]+', value) if v.strip()]
            else:
                parsed = value or None
            if parsed != card[field]:
                card[field] = parsed
                card.setdefault('manual', {})[field] = parsed
                if field == 'company':
                    card['company_evidence'] = None
                changes += 1
    workbook.close()
    save_catalog(root, catalog)
    return {'manual_fields_saved': changes}


def compressed_photo(photo: Path):
    """Preserve existing workbook JPEG bytes; compress native PNGs once."""
    with Image.open(photo) as source:
        if source.format == 'JPEG':
            return photo.read_bytes(), source.width, source.height
        image = source.convert('RGB')
        data = io.BytesIO()
        image.save(data, format='JPEG', quality=92, subsampling=0, optimize=True, dpi=(96, 96))
        return data.getvalue(), image.width, image.height


def insert_fitted_image(sheet, row, column, data, width, height, box_width, box_height,
                        description, url=None, object_position=1):
    scale = min(box_width / width, box_height / height, 1.0)
    options = {'image_data': io.BytesIO(data), 'x_scale': scale, 'y_scale': scale,
               'x_offset': round(6 + (box_width - width * scale) / 2),
               'y_offset': round(6 + (box_height - height * scale) / 2),
               'object_position': object_position, 'description': description}
    if url:
        options['url'] = url
    sheet.insert_image(row, column, 'card.jpg', options)


def gallery_cell(index):
    # Four fixed tiles per row, with a caption and a spacer around each image.
    return 3 + (index // 4) * 3, (index % 4) * 2


def export_library(root: Path, output: Path):
    catalog = load_catalog(root)
    if not catalog['cards']:
        raise ValueError('No imported cards')
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.stem + '.tmp.xlsx')
    book = xlsxwriter.Workbook(temporary, {'strings_to_formulas': False, 'strings_to_urls': False})
    sheet = book.add_worksheet('Contacts')
    gallery = book.add_worksheet('Originals')
    review = book.add_worksheet('Review')
    source_sheet = book.add_worksheet('Sources')
    header = book.add_format({'bold': True, 'bg_color': '#17365D', 'font_color': 'white', 'text_wrap': True})
    normal = book.add_format({'text_wrap': True, 'valign': 'vcenter', 'font_size': 10})
    yellow = book.add_format({'text_wrap': True, 'valign': 'vcenter', 'bg_color': '#FFF2CC', 'font_size': 10})
    text = book.add_format({'text_wrap': True, 'valign': 'vcenter', 'num_format': '@', 'font_size': 10})
    link = book.add_format({'font_color': '#0563C1', 'underline': True, 'valign': 'vcenter', 'text_wrap': True})
    caption = book.add_format({'bold': True, 'text_wrap': True, 'valign': 'vcenter',
                               'bg_color': '#EAF0F7', 'font_size': 11})
    tile = book.add_format({'border': 1, 'border_color': '#DCE4ED', 'bg_color': '#FFFFFF'})
    hint = book.add_format({'font_color': '#596579', 'font_size': 10})
    sheet.write_row(0, 0, HEADERS, header)
    sheet.set_row(0, 30)
    for col, width in enumerate([30, 30, 43, 39, 30, 52, 25]):
        sheet.set_column(col, col, width)
    sheet.set_column(7, 7, 32, None, {'hidden': True})
    sheet.freeze_panes(1, 0)
    sheet.set_zoom(80)
    sheet.hide_gridlines(2)
    gallery.merge_range(0, 0, 0, 6, 'Original business cards', header)
    gallery.merge_range(1, 0, 1, 6,
                        'Four cards per row. Fixed positions for links from the contact table.', hint)
    gallery.set_row(0, 30)
    gallery.set_row(1, 24)
    gallery.freeze_panes(2, 0)
    gallery.hide_gridlines(2)
    gallery.set_zoom(75)
    for col in range(7):
        gallery.set_column_pixels(col, col, 462 if col % 2 == 0 else 18)
    issues = []
    for row, card in enumerate(catalog['cards'], 1):
        sheet.set_row(row, 152)
        flagged = {FIELD_NAMES.get(flag['field'], flag['field']) for flag in card['issues']}
        cells = [card['company'], card['full_name'], card['position'],
                 '\n'.join(card['emails']), '\n'.join(card['phones'])]
        for col, value in enumerate(cells):
            style = yellow if HEADERS[col] in flagged else text if col in (3, 4) else normal
            if value:
                sheet.write_string(row, col, value, style)
            else:
                sheet.write_blank(row, col, None, style)
        sheet.write_string(row, 7, card['id'])
        photo = root / card['photo']
        if not photo.is_file():
            raise FileNotFoundError(photo)
        data, width, height = compressed_photo(photo)
        image_row, image_col = gallery_cell(row - 1)
        target_cell = xl_rowcol_to_cell(image_row, image_col)
        target = "internal:'Originals'!" + target_cell
        sheet.write_url(row, 6, target, link, 'Originals!' + target_cell)
        # A small preview in the table and a larger stable copy on the gallery.
        # XlsxWriter deduplicates identical image bytes, storing one JPEG for both.
        description = card['full_name'] or card['id']
        insert_fitted_image(sheet, row, 5, data, width, height, 340, 188, description, url=target)
        gallery.set_row(image_row - 1, 34)
        gallery.set_row(image_row, 240)
        gallery.set_row(image_row + 1, 14)
        gallery.write_string(image_row - 1, image_col,
                             f'#{row:03} · {description}', caption)
        gallery.write_blank(image_row, image_col, None, tile)
        gallery.write_comment(image_row, image_col, card['id'])
        insert_fitted_image(gallery, image_row, image_col, data, width, height, 450, 308,
                            description, object_position=3)
        evidence = card.get('company_evidence')
        if evidence:
            sheet.write_comment(row, 0, f'Company from domain {evidence["domain"]}: {evidence["url"]}')
        for flag in card['issues']:
            field = FIELD_NAMES.get(flag['field'], flag['field'])
            issues.append([card['id'], row + 1, field, flag.get('visible_fragment', ''), flag.get('reason', '')])
    sheet.autofilter(0, 0, len(catalog['cards']), 7)
    review_headers = ['ID', 'Row', 'Field', 'Visible fragment', 'Reason']
    review.write_row(0, 0, review_headers, header)
    for col, width in enumerate([32, 12, 20, 45, 80]):
        review.set_column(col, col, width)
    for row, issue in enumerate(issues, 1):
        review.write_row(row, 0, issue, normal)
        review.set_row(row, 45)
    review.freeze_panes(1, 0)
    review.autofilter(0, 0, len(issues), 4)
    source_sheet.write_row(0, 0, ['ID', 'Video', 'Card in video', 'Time, seconds', 'SHA256'], header)
    source_sheet.set_column(0, 0, 32)
    source_sheet.set_column(1, 1, 30)
    source_sheet.set_column(2, 3, 22)
    source_sheet.set_column(4, 4, 66)
    row = 1
    for card in catalog['cards']:
        for origin in card['origins']:
            source_sheet.write_row(row, 0, [card['id'], origin['video'], origin['card_id'],
                                           origin['time_sec'], origin['video_sha256']], normal)
            row += 1
    source_sheet.freeze_panes(1, 0)
    book.close()
    import xml.etree.ElementTree as ET
    with ZipFile(temporary) as z:
        if z.testzip():
            raise RuntimeError('Corrupted XLSX')
        ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
              'd': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'}
        for part in (1, 2):
            drawing = ET.fromstring(z.read(f'xl/drawings/drawing{part}.xml'))
            if len(drawing.findall('d:twoCellAnchor', ns)) != len(catalog['cards']):
                raise RuntimeError('Missing table/gallery photos')
        main = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        if len(main.findall('s:hyperlinks/s:hyperlink', ns)) != len(catalog['cards']):
            raise RuntimeError('Missing original links')
        if any(name.startswith('xl/richData/') for name in z.namelist()):
            raise RuntimeError('Unexpected in-cell images')
    temporary.replace(output)
    return {'cards': len(catalog['cards']), 'gallery_columns': 4, 'review_rows': len(issues),
            'xlsx': str(output), 'xlsx_bytes': output.stat().st_size}
