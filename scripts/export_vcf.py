"""Export a business-card XLSX to vCard 3.0 without modifying the workbook."""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


from business_cards.schema import column_index, get_sheet


def text(value):
    return '' if value is None else str(value).strip()


def read_contacts(xlsx: Path):
    book = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        sheet = get_sheet(book, 'contacts')
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows))
        columns = {field: column_index(headers, field)
                   for field in ('company', 'full_name', 'phones')}
        email_col = column_index(headers, 'emails', required=False)
        if email_col is not None:
            columns['emails'] = email_col
        id_col = headers.index('_id') if '_id' in headers else None
        contacts, seen = [], set()
        for row_number, row in enumerate(rows, 2):
            if not any(value is not None and value != '' for value in row):
                continue
            values = {name: text(row[index]) if index < len(row) else ''
                      for name, index in columns.items()}
            if not any(values.values()):
                continue
            stable_id = text(row[id_col]) if id_col is not None and id_col < len(row) else ''
            key = stable_id or f'row:{row_number}'
            if key in seen:
                raise ValueError(f'Duplicate contact key: {key}')
            seen.add(key)
            contacts.append({'key': key, 'row': row_number,
                             'source_name': values['full_name'],
                             'company': values['company'],
                             'raw_emails': values.get('emails', ''),
                             'raw_phones': values['phones']})
        if not contacts:
            raise ValueError('Workbook contains no contacts')
        review = []
        review_sheet = get_sheet(book, 'review', required=False)
        if review_sheet is not None:
            for row in review_sheet.iter_rows(min_row=2, values_only=True):
                if len(row) >= 5 and any(row):
                    review.append({'id': text(row[0]), 'row': row[1],
                                   'field': text(row[2]), 'fragment': text(row[3]),
                                   'reason': text(row[4])})
        return contacts, review
    finally:
        book.close()


def prepare_names(xlsx: Path, output: Path):
    if output.resolve() == xlsx.resolve():
        raise ValueError('Names output must differ from the source workbook')
    contacts, _ = read_contacts(xlsx)
    names = {c['key']: {'source_name': c['source_name'],
                        'first_name': None, 'last_name': None} for c in contacts}
    # Never erase name decisions from a previous preparation.
    with output.open('x', encoding='utf-8') as handle:
        json.dump(names, handle, ensure_ascii=False, indent=2)
        handle.write('\n')
    return {'contacts': len(contacts), 'names': str(output)}


def parse_phone(raw: str):
    # Extension labels describe notation, not the country or type of the number.
    match = re.search(r'\s*,?\s*(?:ext(?:ension)?\.?|x|#|доб(?:авочный)?\.?|вн(?:утр)?\.?)'
                      r'\s*[:.]?\s*([\d\s()\-–—]+)$', raw.strip(), flags=re.I)
    base = raw[:match.start()].strip() if match else raw.strip()
    if base.startswith('+'):
        base = base.replace('(0)', '')
    if not re.fullmatch(r'\+?[\d\s().\-–—]+', base):
        raise ValueError('Unrecognized or unreadable phone notation')
    digits = re.sub(r'\D', '', base)
    if not 3 <= len(digits) <= 15:
        raise ValueError('Invalid phone length')
    # Preserve explicit country codes and national/local numbers as supplied.
    base = ('+' if base.startswith('+') else '') + digits
    extension = re.sub(r'\D', '', match.group(1)) if match else ''
    return {'source': raw, 'value': base + (',,' + extension if extension else '')}


def escape(value: str):
    return (value.replace('\\', '\\\\').replace('\r\n', '\n')
            .replace('\r', '\n').replace('\n', '\\n')
            .replace(';', '\\;').replace(',', '\\,'))


def fold(line: str):
    chunks, current = [], ''
    for char in line:
        if len((current + char).encode('utf-8')) > 75:
            chunks.append(current)
            current = ' '
        current += char
    chunks.append(current)
    return '\r\n'.join(chunks)


def export_vcf(xlsx: Path, names_path: Path, output: Path,
               report_path: Path | None = None):
    paths = [xlsx.resolve(), names_path.resolve(), output.resolve()]
    if report_path is not None:
        paths.append(report_path.resolve())
    if len(set(paths)) != len(paths):
        raise ValueError('Workbook, names, VCF and report paths must differ')
    contacts, review = read_contacts(xlsx)
    names = json.loads(names_path.read_text(encoding='utf-8'))
    if not isinstance(names, dict) or set(names) != {c['key'] for c in contacts}:
        raise ValueError('Name mapping must match every workbook contact exactly')
    cards, records, skipped = [], [], []
    for contact in contacts:
        key = contact['key']
        name = names[key]
        if not isinstance(name, dict) or name.get('source_name') != contact['source_name']:
            raise ValueError(f'{key}: name mapping belongs to a different source name')
        first, last = name.get('first_name'), name.get('last_name')
        if not isinstance(first, str) or not isinstance(last, str):
            raise ValueError(f'{key}: fill first_name and last_name; null is unfinished')
        first, last = first.strip(), last.strip()
        if not first and not last and not contact['company']:
            raise ValueError(f'{key}: contact has neither a name nor a company')
        phones = []
        for raw in re.split(r'[\r\n;]+', contact['raw_phones']):
            if not raw.strip():
                continue
            try:
                phones.append(parse_phone(raw.strip()))
            except ValueError as error:
                skipped.append({'key': key, 'row': contact['row'],
                                'source': raw.strip(), 'reason': str(error)})
        selected = list({p['value']: p for p in phones}.values())
        emails = list(dict.fromkeys(email.strip()
                                    for email in re.split(r'[\r\n;,]+', contact['raw_emails'])
                                    if email.strip()))
        fields = ['BEGIN:VCARD', 'VERSION:3.0',
                  'N:' + ';'.join((escape(last), escape(first), '', '', '')),
                  'FN:' + escape(' '.join(p for p in (first, last) if p) or contact['company']),
                  'ORG:' + escape(contact['company'])]
        for email in emails:
            fields.append('EMAIL;TYPE=INTERNET:' + escape(email))
        for phone in selected:
            fields.append('TEL;TYPE=VOICE:' + escape(phone['value']))
        fields.append('END:VCARD')
        cards.append('\r\n'.join(fold(field) for field in fields) + '\r\n')
        records.append({**contact, 'first_name': first, 'last_name': last,
                        'selected_phones': selected,
                        'emails': emails})
    report = {
        'contacts': len(records),
        'with_phone': sum(bool(r['selected_phones']) for r in records),
        'without_phone': sum(not r['selected_phones'] for r in records),
        'telephone_fields': sum(len(r['selected_phones']) for r in records),
        'email_fields': sum(len(r['emails']) for r in records),
        'skipped_phone_count': len(skipped), 'skipped_phones': skipped,
        'source_review': review, 'records': records,
    }
    # All records have been validated before writing either artifact.
    output.write_bytes(''.join(cards).encode('utf-8'))
    if report_path is not None:
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    summary = {k: v for k, v in report.items()
               if k not in ('records', 'source_review', 'skipped_phones')}
    summary.update({'output': str(output), 'report': str(report_path) if report_path else None})
    if skipped and report_path is None:
        summary['skipped_phones'] = skipped
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest='command', required=True)
    prepare = subparsers.add_parser('prepare-names', help='Create a name mapping for review in chat')
    prepare.add_argument('xlsx', type=Path)
    prepare.add_argument('--output', type=Path, required=True)
    export = subparsers.add_parser('export', help='Export reviewed names and selected phones')
    export.add_argument('xlsx', type=Path)
    export.add_argument('--names', type=Path, required=True)
    export.add_argument('--output', type=Path, required=True)
    export.add_argument('--report', type=Path)
    args = parser.parse_args()
    try:
        if args.command == 'prepare-names':
            result = prepare_names(args.xlsx, args.output)
        else:
            result = export_vcf(args.xlsx, args.names, args.output, args.report)
    except (ValueError, KeyError, OSError) as error:
        parser.exit(1, f'Error: {error}\n')
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
